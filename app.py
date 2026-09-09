# OBS! Kör den här filen i terminalen med kommandot:
#     streamlit run app.py

import streamlit as st
import requests
from openpyxl import load_workbook
from io import BytesIO

FILNAMN = "Allsvenskan_Tipstävling_2026.xlsx"


def hämta_och_uppdatera_excel():

    # === Ny GraphQL-adress ===
    url = "https://gql.sportomedia.se/graphql"

    # === GraphQL-fråga ===
    query = """
    query matchesForLeague(
      $configLeagueName: String!
      $configSeasonStartYear: Int!
    ) {
      matchesForLeague(
        configLeagueName: $configLeagueName
        configSeasonStartYear: $configSeasonStartYear
      ) {
        matches {
          id
          startDate
          homeTeamName
          visitingTeamName
          homeTeamScore
          visitingTeamScore
          status
          extendedStatus
          period
          round
        }
      }
    }
    """

    # === Variabler ===
    variables = {
        "configLeagueName": "allsvenskan",
        "configSeasonStartYear": 2026
    }

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "Accept": "application/json",
        "Content-Type": "application/json",
        "Origin": "https://allsvenskan.se",
        "Referer": "https://allsvenskan.se/statistik/tabell"
    }

    # === Hämta data från GraphQL ===
    try:
        response = requests.post(
            url,
            json={
                "query": query,
                "variables": variables
            },
            headers=headers,
            timeout=20
        )

        response.raise_for_status()

        result = response.json()

        if "errors" in result:
            st.error(f"❌ GraphQL-fel: {result['errors']}")
            return None

        matcher = result["data"]["matchesForLeague"]["matches"]

    except requests.RequestException as e:
        st.error(f"❌ Kunde inte hämta data från GraphQL: {e}")
        return None

    except (KeyError, TypeError, ValueError) as e:
        st.error(f"❌ Kunde inte läsa data från GraphQL: {e}")
        return None

    # === Bygg tabellen från färdigspelade matcher ===
    lagdata = {}

    for match in matcher:

        # Räkna ENDAST färdigspelade matcher
        if match.get("status") != "FINISHED":
            continue

        hemma = match.get("homeTeamName")
        borta = match.get("visitingTeamName")

        hemma_mal = match.get("homeTeamScore")
        borta_mal = match.get("visitingTeamScore")

        # Om något resultat saknas, hoppa över matchen
        if hemma_mal is None or borta_mal is None:
            continue

        # Skapa lagen första gången vi ser dem
        for lag in (hemma, borta):
            if lag not in lagdata:
                lagdata[lag] = {
                    "S": 0,
                    "V": 0,
                    "O": 0,
                    "F": 0,
                    "GM": 0,
                    "IM": 0
                }

        # === Hemmalag ===
        lagdata[hemma]["S"] += 1
        lagdata[hemma]["GM"] += hemma_mal
        lagdata[hemma]["IM"] += borta_mal

        # === Bortalag ===
        lagdata[borta]["S"] += 1
        lagdata[borta]["GM"] += borta_mal
        lagdata[borta]["IM"] += hemma_mal

        # === Resultat ===
        if hemma_mal > borta_mal:
            lagdata[hemma]["V"] += 1
            lagdata[borta]["F"] += 1

        elif hemma_mal < borta_mal:
            lagdata[borta]["V"] += 1
            lagdata[hemma]["F"] += 1

        else:
            lagdata[hemma]["O"] += 1
            lagdata[borta]["O"] += 1

    # === Målskillnad och poäng ===
    for lag, stats in lagdata.items():
        stats["MS"] = stats["GM"] - stats["IM"]
        stats["P"] = stats["V"] * 3 + stats["O"]

    # === Sortera tabellen ===
    sorterade_lag = sorted(
        lagdata.items(),
        key=lambda x: (
            -x[1]["P"],
            -x[1]["MS"],
            -x[1]["GM"]
        )
    )

    # === Öppna Excel-filen ===
    try:
        wb = load_workbook(FILNAMN)
        ws = wb["Allsvenskan"]

    except Exception as e:
        st.error(f"❌ Kunde inte öppna Excel-filen: {e}")
        return None

    # === Skriv tabellen till Excel ===
    for i, (lag_namn, stats) in enumerate(sorterade_lag[:16], start=2):

        ws.cell(row=i, column=1).value = i - 1
        ws.cell(row=i, column=2).value = lag_namn
        ws.cell(row=i, column=3).value = stats["S"]
        ws.cell(row=i, column=4).value = stats["V"]
        ws.cell(row=i, column=5).value = stats["O"]
        ws.cell(row=i, column=6).value = stats["F"]
        ws.cell(row=i, column=7).value = stats["GM"]
        ws.cell(row=i, column=8).value = stats["IM"]
        ws.cell(row=i, column=9).value = stats["MS"]
        ws.cell(row=i, column=10).value = stats["P"]

    # === Spara Excel-filen i minnet ===
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return stream


st.title("Allsvenskan Tipstävling 2026")


# === Uppdatera-knappen först ===
if st.button("🔄 Uppdatera Excel-fil med senaste data"):

    excel_stream = hämta_och_uppdatera_excel()

    if excel_stream:
        st.success("✅ Excel-filen är uppdaterad! Du kan ladda ner den nedan.")
        st.session_state["excel_stream"] = excel_stream

    else:
        st.error("❌ Kunde inte uppdatera Excel-filen.")


# === Info om knappens funktion om fil ej laddats ner ännu ===
if "excel_stream" not in st.session_state:
    st.info("Tryck på knappen ovan för att hämta senaste data och ladda ner Excel-filen.")


# === Ladda ner-knappen direkt efter ===
if "excel_stream" in st.session_state:

    st.download_button(
        label="📥 Ladda ner uppdaterad Excel-fil",
        data=st.session_state["excel_stream"],
        file_name=FILNAMN,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# === Viktig info om redigering ===
st.markdown("### ℹ️ När du öppnar Excel-filen")

st.markdown("""
👉 **Klicka på `Aktivera redigering`** upptill i fönstret om du använder Excel.  
Det behövs för att formlerna och resultaten ska visas korrekt.
""")


# === Hjälp för de utan Excel ===
with st.expander("📂 Har du inte Excel installerat?"):

    st.markdown("""
Du kan fortfarande öppna och använda filen:

**🔹 Excel Online:**  
Gå till [office.com/launch/excel](https://office.com/launch/excel) – kräver Microsoft-konto.

**🔹 Google Kalkylark:**  
1. Gå till [Google Drive](https://drive.google.com)  
2. Ladda upp Excel-filen  
3. Högerklicka → **Öppna med → Google Kalkylark**  
4. Observera att vissa formler kan visas annorlunda än i Excel.
""")
